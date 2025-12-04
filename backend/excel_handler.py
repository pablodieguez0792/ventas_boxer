"""
Manejo de archivos Excel para carga y descarga masiva de productos
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from io import BytesIO
from datetime import datetime

class ExcelHandler:
    """Clase para manejar operaciones con Excel"""
    
    @staticmethod
    def create_product_template():
        """Crear plantilla Excel para carga masiva de productos"""
        
        # Definir columnas según API de Tienda Nube
        columns = [
            'nombre',
            'descripcion',
            'precio',
            'precio_promocional',
            'stock',
            'sku',
            'peso_kg',
            'ancho_cm',
            'alto_cm',
            'profundidad_cm',
            'publicado',
            'envio_gratis',
            'marca',
            'categoria',
            'url_imagen_1',
            'url_imagen_2',
            'url_imagen_3',
            'url_imagen_4',
            'url_imagen_5',
            'seo_titulo',
            'seo_descripcion',
            'tags',
            'codigo_barras',
            'mpn',
            'age_group',
            'gender',
        ]
        
        # Crear workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Productos"
        
        # Agregar encabezados con estilo
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for col_num, column_name in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = column_name
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Agregar fila de ejemplo
        example_row = [
            'Producto de Ejemplo',  # nombre
            'Descripción detallada del producto',  # descripcion
            '10000',  # precio
            '8500',  # precio_promocional
            '50',  # stock
            'SKU-001',  # sku
            '0.5',  # peso_kg
            '10',  # ancho_cm
            '15',  # alto_cm
            '20',  # profundidad_cm
            'SI',  # publicado
            'NO',  # envio_gratis
            'Mi Marca',  # marca
            'Categoría Principal',  # categoria
            'https://ejemplo.com/imagen1.jpg',  # url_imagen_1
            '',  # url_imagen_2
            '',  # url_imagen_3
            '',  # url_imagen_4
            '',  # url_imagen_5
            'Título SEO del Producto',  # seo_titulo
            'Descripción SEO del producto',  # seo_descripcion
            'tag1,tag2,tag3',  # tags
            '7891234567890',  # codigo_barras
            'MPN-001',  # mpn
            'adult',  # age_group
            'unisex',  # gender
        ]
        
        for col_num, value in enumerate(example_row, 1):
            ws.cell(row=2, column=col_num, value=value)
        
        # Ajustar ancho de columnas
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Agregar hoja de instrucciones
        ws_instructions = wb.create_sheet("Instrucciones")
        instructions = [
            ["INSTRUCCIONES PARA CARGA MASIVA DE PRODUCTOS"],
            [""],
            ["1. Complete la hoja 'Productos' con los datos de sus productos"],
            ["2. Los campos obligatorios son: nombre, precio, stock"],
            ["3. Formato de campos:"],
            ["   - precio/precio_promocional: números sin símbolos (ej: 10000)"],
            ["   - peso_kg: peso en kilogramos con punto decimal (ej: 0.5)"],
            ["   - dimensiones: en centímetros (ej: 10, 15, 20)"],
            ["   - publicado/envio_gratis: SI o NO"],
            ["   - tags: separados por comas (ej: tag1,tag2,tag3)"],
            ["   - url_imagen_X: URLs completas de las imágenes"],
            [""],
            ["4. Puede agregar tantas filas como productos desee"],
            ["5. Elimine la fila de ejemplo antes de subir el archivo"],
            ["6. Guarde el archivo y súbalo en la sección 'Agregar Producto'"],
            [""],
            ["CAMPOS OPCIONALES:"],
            ["- precio_promocional: precio en oferta"],
            ["- peso_kg, ancho_cm, alto_cm, profundidad_cm: dimensiones del producto"],
            ["- marca, categoria: información de clasificación"],
            ["- seo_titulo, seo_descripcion: optimización para buscadores"],
            ["- codigo_barras, mpn: códigos de identificación"],
            ["- age_group: newborn, infant, toddler, kids, adult"],
            ["- gender: female, male, unisex"],
        ]
        
        for row_num, instruction in enumerate(instructions, 1):
            ws_instructions.cell(row=row_num, column=1, value=instruction[0])
        
        ws_instructions.column_dimensions['A'].width = 80
        
        # Guardar en BytesIO
        excel_file = BytesIO()
        wb.save(excel_file)
        excel_file.seek(0)
        
        return excel_file
    
    @staticmethod
    def parse_products_from_excel(file_content):
        """Parsear productos desde archivo Excel"""
        
        try:
            df = pd.read_excel(file_content, sheet_name='Productos')
            
            products = []
            
            for index, row in df.iterrows():
                # Saltar filas vacías o de ejemplo
                if pd.isna(row.get('nombre')) or row.get('nombre') == 'Producto de Ejemplo':
                    continue
                
                # Construir objeto de producto
                product = {
                    'name': {'es': str(row.get('nombre', ''))},
                    'description': {'es': str(row.get('descripcion', ''))},
                    'published': str(row.get('publicado', 'SI')).upper() == 'SI',
                    'free_shipping': str(row.get('envio_gratis', 'NO')).upper() == 'SI',
                }
                
                # Variante principal
                variant = {
                    'price': str(row.get('precio', 0)),
                    'stock': int(row.get('stock', 0)) if not pd.isna(row.get('stock')) else 0,
                }
                
                # Campos opcionales de variante
                if not pd.isna(row.get('precio_promocional')):
                    variant['promotional_price'] = str(row.get('precio_promocional'))
                
                if not pd.isna(row.get('sku')):
                    variant['sku'] = str(row.get('sku'))
                
                if not pd.isna(row.get('peso_kg')):
                    variant['weight'] = str(row.get('peso_kg'))
                
                if not pd.isna(row.get('ancho_cm')):
                    variant['width'] = str(row.get('ancho_cm'))
                
                if not pd.isna(row.get('alto_cm')):
                    variant['height'] = str(row.get('alto_cm'))
                
                if not pd.isna(row.get('profundidad_cm')):
                    variant['depth'] = str(row.get('profundidad_cm'))
                
                if not pd.isna(row.get('codigo_barras')):
                    variant['barcode'] = str(row.get('codigo_barras'))
                
                if not pd.isna(row.get('mpn')):
                    variant['mpn'] = str(row.get('mpn'))
                
                product['variants'] = [variant]
                
                # Imágenes
                images = []
                for i in range(1, 6):
                    url = row.get(f'url_imagen_{i}')
                    if not pd.isna(url) and str(url).strip():
                        images.append({'src': str(url).strip()})
                
                if images:
                    product['images'] = images
                
                # Campos opcionales de producto
                if not pd.isna(row.get('marca')):
                    product['brand'] = str(row.get('marca'))
                
                if not pd.isna(row.get('seo_titulo')):
                    product['seo_title'] = {'es': str(row.get('seo_titulo'))[:70]}
                
                if not pd.isna(row.get('seo_descripcion')):
                    product['seo_description'] = {'es': str(row.get('seo_descripcion'))[:160]}
                
                if not pd.isna(row.get('tags')):
                    tags_str = str(row.get('tags'))
                    product['tags'] = [tag.strip() for tag in tags_str.split(',') if tag.strip()]
                
                products.append(product)
            
            return products
        
        except Exception as e:
            raise Exception(f"Error al parsear Excel: {str(e)}")
    
    @staticmethod
    def export_products_to_excel(products):
        """Exportar productos a Excel"""
        
        data = []
        
        for product in products:
            variant = product.get('variants', [{}])[0]
            images = product.get('images', [])
            
            row = {
                'id': product.get('id'),
                'nombre': product.get('name', {}).get('es', ''),
                'descripcion': product.get('description', {}).get('es', ''),
                'precio': variant.get('price', ''),
                'precio_promocional': variant.get('promotional_price', ''),
                'stock': variant.get('stock', 0),
                'sku': variant.get('sku', ''),
                'peso_kg': variant.get('weight', ''),
                'ancho_cm': variant.get('width', ''),
                'alto_cm': variant.get('height', ''),
                'profundidad_cm': variant.get('depth', ''),
                'publicado': 'SI' if product.get('published') else 'NO',
                'envio_gratis': 'SI' if product.get('free_shipping') else 'NO',
                'marca': product.get('brand', ''),
                'seo_titulo': product.get('seo_title', {}).get('es', ''),
                'seo_descripcion': product.get('seo_description', {}).get('es', ''),
                'tags': ','.join(product.get('tags', [])),
                'codigo_barras': variant.get('barcode', ''),
                'mpn': variant.get('mpn', ''),
                'variant_id': variant.get('id', ''),
            }
            
            # Agregar URLs de imágenes
            for i in range(5):
                if i < len(images):
                    row[f'url_imagen_{i+1}'] = images[i].get('src', '')
                else:
                    row[f'url_imagen_{i+1}'] = ''
            
            data.append(row)
        
        df = pd.DataFrame(data)
        
        # Crear archivo Excel
        excel_file = BytesIO()
        with pd.ExcelWriter(excel_file, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Productos', index=False)
            
            # Obtener workbook para aplicar estilos
            workbook = writer.book
            worksheet = writer.sheets['Productos']
            
            # Estilo de encabezados
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
            
            # Ajustar ancho de columnas
            for column in worksheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                worksheet.column_dimensions[column_letter].width = adjusted_width
        
        excel_file.seek(0)
        return excel_file
